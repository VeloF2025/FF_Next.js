"""
Excel Generator for Reconciliation Reports

Creates multi-sheet Excel files with formatted tables and styling
for daily reconciliation reports.
"""

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
import pandas as pd

from lib.reconciliation.models import ReconciliationReport

logger = logging.getLogger(__name__)


class ReconciliationExcelGenerator:
    """
    Generates formatted Excel reports from ReconciliationReport data.

    Creates 5 sheets:
    - Summary: Key metrics and statistics
    - WA_Submissions: All WhatsApp Monitor submissions
    - OES_Activations: All OES (Nokia_Exp) activations
    - WA_Only_Gaps: Submissions without corresponding activation
    - OES_Only_Gaps: Activations without corresponding submission
    """

    # Styling constants
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
    TITLE_FONT = Font(bold=True, size=14)
    METRIC_FONT = Font(bold=True, size=11)
    THIN_BORDER = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Gap highlighting
    GAP_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    GOOD_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    WARNING_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

    def __init__(self, output_dir: str = "data/reconciliation"):
        """
        Initialize generator with output directory.

        Args:
            output_dir: Directory to save generated reports
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(
        self,
        report: ReconciliationReport,
        filename: Optional[str] = None
    ) -> Path:
        """
        Generate Excel report from reconciliation data.

        Args:
            report: ReconciliationReport with all data
            filename: Optional custom filename (default: project_reconciliation_date.xlsx)

        Returns:
            Path to generated file
        """
        # Generate filename if not provided
        if not filename:
            filename = (
                f"{report.summary.project}_reconciliation_"
                f"{report.summary.report_date.strftime('%Y-%m-%d')}.xlsx"
            )

        output_path = self.output_dir / filename

        logger.info(f"Generating reconciliation report: {output_path}")

        # Create workbook
        wb = Workbook()

        # Remove default sheet
        wb.remove(wb.active)

        # Create sheets
        self._create_summary_sheet(wb, report)
        self._create_wa_submissions_sheet(wb, report)
        self._create_oes_activations_sheet(wb, report)
        self._create_wa_only_sheet(wb, report)
        self._create_oes_only_sheet(wb, report)

        # Save workbook
        wb.save(output_path)

        logger.info(f"Report saved: {output_path}")

        return output_path

    def _create_summary_sheet(self, wb: Workbook, report: ReconciliationReport) -> None:
        """Create Summary sheet with key metrics."""
        ws = wb.create_sheet("Summary")

        summary = report.summary

        # Title
        ws["A1"] = f"Reconciliation Report: {summary.project_display_name}"
        ws["A1"].font = self.TITLE_FONT
        ws.merge_cells("A1:C1")

        # Report metadata
        ws["A3"] = "Report Date:"
        ws["B3"] = summary.report_date.strftime("%Y-%m-%d")
        ws["A4"] = "Generated At:"
        ws["B4"] = summary.generated_at.strftime("%Y-%m-%d %H:%M:%S")
        ws["A5"] = "Project:"
        ws["B5"] = summary.project_display_name

        # Data source info
        ws["A7"] = "Data Sources"
        ws["A7"].font = self.METRIC_FONT
        ws["A8"] = "WA Source:"
        ws["B8"] = summary.wa_data_source
        ws["A9"] = "OES Source:"
        ws["B9"] = summary.oes_data_source

        # Metrics section
        ws["A11"] = "Reconciliation Metrics"
        ws["A11"].font = self.METRIC_FONT

        metrics = [
            ("WhatsApp Submissions", summary.wa_submission_count),
            ("OES Activations", summary.oes_activation_count),
            ("Matched Records", summary.matched_count),
            ("WA Only (Not Activated)", summary.wa_only_count),
            ("OES Only (No QA Photo)", summary.oes_only_count),
        ]

        row = 12
        for label, value in metrics:
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = Font(bold=True)
            row += 1

        # Percentages
        row += 1
        ws[f"A{row}"] = "Match Rate (WA → OES)"
        ws[f"B{row}"] = f"{summary.match_rate:.1f}%"
        ws[f"A{row}"].font = Font(bold=True)

        # Color code match rate
        if summary.match_rate >= 90:
            ws[f"B{row}"].fill = self.GOOD_FILL
        elif summary.match_rate >= 70:
            ws[f"B{row}"].fill = self.WARNING_FILL
        else:
            ws[f"B{row}"].fill = self.GAP_FILL

        row += 1
        ws[f"A{row}"] = "Activation Coverage"
        ws[f"B{row}"] = f"{summary.activation_coverage:.1f}%"
        ws[f"A{row}"].font = Font(bold=True)

        # Adjust column widths
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 20

    def _create_wa_submissions_sheet(
        self,
        wb: Workbook,
        report: ReconciliationReport
    ) -> None:
        """Create WA_Submissions sheet with all WhatsApp data."""
        ws = wb.create_sheet("WA_Submissions")

        if not report.wa_submissions:
            ws["A1"] = "No WhatsApp submissions found"
            return

        # Convert to DataFrame for easier handling
        data = []
        for sub in report.wa_submissions:
            data.append({
                "DR_Number": sub.dr_number,
                "Submission_Date": sub.submission_date.strftime("%Y-%m-%d") if sub.submission_date else "",
                "Verification_Steps": f"{sub.verification_steps}/12",
                "QA_Verified": "Yes" if sub.qa_verified else "No",
                "Contractor": sub.contractor or "",
                "Photo_Count": sub.photo_count,
                "Status": sub.status or "",
                "Notes": sub.notes or ""
            })

        df = pd.DataFrame(data)
        self._write_dataframe_to_sheet(ws, df)

    def _create_oes_activations_sheet(
        self,
        wb: Workbook,
        report: ReconciliationReport
    ) -> None:
        """Create OES_Activations sheet with all OES data."""
        ws = wb.create_sheet("OES_Activations")

        if not report.oes_activations:
            ws["A1"] = "No OES activations found"
            return

        # Convert to DataFrame
        data = []
        for act in report.oes_activations:
            data.append({
                "DR_Number": act.dr_number,
                "OLT": act.olt,
                "Activation_Date": act.activation_date.strftime("%Y-%m-%d") if act.activation_date else "",
                "Status": act.status,
                "Customer_Name": act.customer_name or "",
                "Service_Type": act.service_type or "",
                "ONT_Serial": act.ont_serial or ""
            })

        df = pd.DataFrame(data)
        self._write_dataframe_to_sheet(ws, df)

    def _create_wa_only_sheet(self, wb: Workbook, report: ReconciliationReport) -> None:
        """Create WA_Only_Gaps sheet - submissions without activation."""
        ws = wb.create_sheet("WA_Only_Gaps")

        # Add description
        ws["A1"] = "Drops Done - Not Yet Activated"
        ws["A1"].font = self.TITLE_FONT
        ws.merge_cells("A1:D1")
        ws["A2"] = "These DR numbers have QA photo verification but no OES activation"

        if not report.wa_only_records:
            ws["A4"] = "No gaps found - all submissions are activated!"
            ws["A4"].fill = self.GOOD_FILL
            return

        # Convert to DataFrame
        data = []
        for sub in report.wa_only_records:
            data.append({
                "DR_Number": sub.dr_number,
                "Submission_Date": sub.submission_date.strftime("%Y-%m-%d") if sub.submission_date else "",
                "Verification_Steps": f"{sub.verification_steps}/12",
                "QA_Verified": "Yes" if sub.qa_verified else "No",
                "Contractor": sub.contractor or "",
                "Photo_Count": sub.photo_count,
                "Notes": "Pending activation"
            })

        df = pd.DataFrame(data)
        self._write_dataframe_to_sheet(ws, df, start_row=4)

        # Highlight all gap rows
        for row in range(5, 5 + len(data)):
            for col in range(1, len(df.columns) + 1):
                ws.cell(row=row, column=col).fill = self.GAP_FILL

    def _create_oes_only_sheet(self, wb: Workbook, report: ReconciliationReport) -> None:
        """Create OES_Only_Gaps sheet - activations without QA photos."""
        ws = wb.create_sheet("OES_Only_Gaps")

        # Add description
        ws["A1"] = "Activated Without QA Photo Verification"
        ws["A1"].font = self.TITLE_FONT
        ws.merge_cells("A1:D1")
        ws["A2"] = "These DR numbers are activated but have no WhatsApp Monitor submission"

        if not report.oes_only_records:
            ws["A4"] = "No gaps found - all activations have QA verification!"
            ws["A4"].fill = self.GOOD_FILL
            return

        # Convert to DataFrame
        data = []
        for act in report.oes_only_records:
            data.append({
                "DR_Number": act.dr_number,
                "OLT": act.olt,
                "Activation_Date": act.activation_date.strftime("%Y-%m-%d") if act.activation_date else "",
                "Status": act.status,
                "Customer_Name": act.customer_name or "",
                "Notes": "Missing QA photo"
            })

        df = pd.DataFrame(data)
        self._write_dataframe_to_sheet(ws, df, start_row=4)

        # Highlight all gap rows
        for row in range(5, 5 + len(data)):
            for col in range(1, len(df.columns) + 1):
                ws.cell(row=row, column=col).fill = self.GAP_FILL

    def _write_dataframe_to_sheet(
        self,
        ws,
        df: pd.DataFrame,
        start_row: int = 1
    ) -> None:
        """Write DataFrame to worksheet with formatting."""
        # Write headers
        for col_idx, column in enumerate(df.columns, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=column)
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
            cell.border = self.THIN_BORDER

        # Write data
        for row_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(horizontal="left")

        # Auto-adjust column widths
        for col_idx, column in enumerate(df.columns, 1):
            max_length = len(str(column))
            for row in df.itertuples(index=False):
                value = row[col_idx - 1]
                if value:
                    max_length = max(max_length, len(str(value)))

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width


def generate_reconciliation_excel(
    report: ReconciliationReport,
    output_dir: str = "data/reconciliation",
    filename: Optional[str] = None
) -> Path:
    """
    Convenience function to generate Excel report.

    Args:
        report: ReconciliationReport data
        output_dir: Output directory path
        filename: Optional custom filename

    Returns:
        Path to generated file
    """
    generator = ReconciliationExcelGenerator(output_dir)
    return generator.generate(report, filename)
